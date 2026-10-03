"""Tests for behave_modern_file_report.docx_writer."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx.shared import RGBColor

from behave_modern_file_report.docx_writer import DOCXWriter
from behave_modern_file_report.models import (
    Attachment,
    Background,
    Environment,
    ErrorInfo,
    FeatureSummary,
    ReportOptions,
    RunSummary,
    ScenarioResult,
    Step,
)


def _make_step(
    name: str = "step",
    keyword: str = "Given ",
    status: str = "passed",
    duration: float = 0.01,
    error: ErrorInfo | None = None,
    attachments: list[Attachment] | None = None,
    logs: list[str] | None = None,
) -> Step:
    return Step(
        keyword=keyword,
        name=name,
        status=status,
        duration=duration,
        location="f:10",
        error=error,
        attachments=attachments or [],
        logs=logs or [],
    )


def _make_scenario(
    name: str = "S1",
    status: str = "passed",
    steps: list[Step] | None = None,
    error: ErrorInfo | None = None,
    tags: list[str] | None = None,
    is_outline: bool = False,
    background: Background | None = None,
) -> ScenarioResult:
    return ScenarioResult(
        name=name,
        status=status,
        duration=0.05,
        tags=tags or [],
        location="f:5",
        feature_name="F1",
        is_outline=is_outline,
        steps=steps or [],
        error=error,
        background=background,
    )


def _make_feature(
    name: str = "F1",
    scenarios: list[ScenarioResult] | None = None,
    background: Background | None = None,
    tags: list[str] | None = None,
    description: str = "",
) -> FeatureSummary:
    return FeatureSummary(
        name=name,
        description=description,
        duration=0.1,
        tags=tags or [],
        location="f:1",
        scenarios=scenarios or [],
        background=background,
    )


def _make_run(features: list[FeatureSummary] | None = None, title: str = "Test") -> RunSummary:
    return RunSummary(
        run_id="r1",
        title=title,
        project_name="P",
        start_time="2025-01-01T10:00:00+00:00",
        end_time="2025-01-01T10:00:05+00:00",
        duration=5.0,
        features=features or [],
        environment=Environment(python_version="3.14", platform="linux"),
    )


def _write(run: RunSummary, tmp_path: Path, opts: ReportOptions | None = None) -> str:
    path = tmp_path / "report.docx"
    DOCXWriter(opts or ReportOptions()).write(run, path)
    return str(path)


def _read_docx(path: str) -> Any:
    from docx import Document

    return Document(path)


def _all_text(doc: Any) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    parts.append(para.text)
    return "\n".join(parts)


# --- Basic generation ---


def test_generates_valid_docx(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario()])]), tmp_path)
    doc = _read_docx(path)
    assert len(doc.paragraphs) > 0


def test_cover_title(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature()]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Test" in text


def test_cover_project_name(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature()]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "P" in text


def test_cover_metadata(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature()]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Started" in text
    assert "2025-01-01" in text


def test_summary_section(tmp_path: Path) -> None:
    feat = _make_feature(
        scenarios=[_make_scenario(status="passed"), _make_scenario(status="failed")],
    )
    path = _write(_make_run([feat]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Executive Summary" in text
    assert "Scenarios" in text


def test_feature_heading(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature(name="Login")]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Login" in text


def test_scenario_heading(tmp_path: Path) -> None:
    feat = _make_feature(scenarios=[_make_scenario(name="Login works")])
    path = _write(_make_run([feat]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Login works" in text


def test_step_in_table(tmp_path: Path) -> None:
    step = _make_step(name="user on page")
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "user on page" in text


def test_status_labels(tmp_path: Path) -> None:
    feat = _make_feature(
        scenarios=[_make_scenario(status="passed"), _make_scenario(status="failed")],
    )
    path = _write(_make_run([feat]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "PASSED" in text
    assert "FAILED" in text


def test_error_block(tmp_path: Path) -> None:
    err = ErrorInfo(message="boom", traceback="Traceback:\n  File x", exception_type="ValueError")
    feat = _make_feature(scenarios=[_make_scenario(status="failed", error=err)])
    path = _write(_make_run([feat]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Failure" in text
    assert "boom" in text
    assert "ValueError" in text


def test_error_without_traceback(tmp_path: Path) -> None:
    err = ErrorInfo(message="no tb", traceback="", exception_type="RuntimeError")
    feat = _make_feature(scenarios=[_make_scenario(status="failed", error=err)])
    path = _write(_make_run([feat]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "no tb" in text


def test_error_without_exception_type(tmp_path: Path) -> None:
    err = ErrorInfo(message="unknown", traceback="", exception_type="")
    feat = _make_feature(scenarios=[_make_scenario(status="failed", error=err)])
    path = _write(_make_run([feat]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "unknown" in text


def test_feature_tags(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature(tags=["auth", "smoke"])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "auth" in text
    assert "smoke" in text


def test_scenario_tags(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(tags=["fast"])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "fast" in text


def test_feature_description(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature(description="A test feature")]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "A test feature" in text


def test_outline_tag(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(is_outline=True)])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "OUTLINE" in text


def test_rule_name(tmp_path: Path) -> None:
    scn = _make_scenario()
    scn.rule_name = "Auth rule"
    path = _write(_make_run([_make_feature(scenarios=[scn])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Auth rule" in text


def test_background(tmp_path: Path) -> None:
    bg = Background(name="Setup", steps=[_make_step(name="app running")])
    path = _write(_make_run([_make_feature(background=bg, scenarios=[_make_scenario()])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Background" in text
    assert "Setup" in text
    assert "app running" in text


def test_scenario_background(tmp_path: Path) -> None:
    bg = Background(name="Scn bg", steps=[_make_step(name="setup")])
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(background=bg)])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Scn bg" in text


def test_attachments_listed(tmp_path: Path) -> None:
    step = _make_step(name="step", attachments=[Attachment(name="shot.png", mime_type="image/png")])
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "shot.png" in text


def test_step_logs(tmp_path: Path) -> None:
    step = _make_step(name="step", logs=["navigated", "clicked"])
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "navigated" in text
    assert "clicked" in text


def test_step_error_in_table(tmp_path: Path) -> None:
    err = ErrorInfo(message="step err", traceback="", exception_type="Err")
    step = _make_step(name="bad", status="failed", error=err)
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "step err" in text


def test_environment_section(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature()]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Environment" in text
    assert "3.14" in text
    assert "linux" in text


def test_toc_placeholder(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature()]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Table of Contents" in text


def test_multiple_features(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature(name="A"), _make_feature(name="B")]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "A" in text
    assert "B" in text


def test_empty_run(tmp_path: Path) -> None:
    path = _write(_make_run([]), tmp_path)
    doc = _read_docx(path)
    assert len(doc.paragraphs) > 0


def test_feature_without_tags_description_location(tmp_path: Path) -> None:
    feat = _make_feature(name="Minimal", tags=[], description="")
    feat.location = ""
    path = _write(_make_run([feat]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Minimal" in text


def test_scenario_without_tags_rule_background(tmp_path: Path) -> None:
    scn = _make_scenario(name="Min", tags=[])
    scn.rule_name = ""
    scn.background = None
    path = _write(_make_run([_make_feature(scenarios=[scn])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Min" in text


def test_skipped_status(tmp_path: Path) -> None:
    feat = _make_feature(scenarios=[_make_scenario(status="skipped")])
    path = _write(_make_run([feat]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "SKIPPED" in text


def test_undefined_status(tmp_path: Path) -> None:
    feat = _make_feature(scenarios=[_make_scenario(status="undefined")])
    path = _write(_make_run([feat]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "UNDEFINED" in text


def test_step_table_empty_steps(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[])])]), tmp_path)
    doc = _read_docx(path)
    assert len(doc.paragraphs) > 0


def test_feature_summary_table(tmp_path: Path) -> None:
    feat = _make_feature(
        scenarios=[_make_scenario(status="passed"), _make_scenario(status="failed")],
    )
    path = _write(_make_run([feat]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Pass rate" in text


def test_progress_bar(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario()])]), tmp_path)
    doc = _read_docx(path)
    assert len(doc.tables) > 0


def test_custom_title(tmp_path: Path) -> None:
    path = _write(_make_run([_make_feature()], title="Custom"), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Custom" in text


def test_cover_without_project(tmp_path: Path) -> None:
    run = _make_run([_make_feature()])
    run.project_name = ""
    path = _write(run, tmp_path)
    doc = _read_docx(path)
    assert len(doc.paragraphs) > 0


def test_environment_empty_fields(tmp_path: Path) -> None:
    run = _make_run([_make_feature()])
    run.environment = Environment()
    path = _write(run, tmp_path)
    text = _all_text(_read_docx(path))
    assert "Environment" in text


def test_scenario_without_location(tmp_path: Path) -> None:
    scn = _make_scenario(name="NoLoc")
    scn.location = ""
    path = _write(_make_run([_make_feature(scenarios=[scn])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "NoLoc" in text


def test_step_with_unknown_status(tmp_path: Path) -> None:
    step = _make_step(name="weird", status="untested")
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "weird" in text


def test_status_badge_unknown_status(tmp_path: Path) -> None:
    feat = _make_feature(scenarios=[_make_scenario(status="untested")])
    path = _write(_make_run([feat]), tmp_path)
    doc = _read_docx(path)
    assert len(doc.tables) > 0


# ---------------------------------------------------------------------------
# Attachment rendering
# ---------------------------------------------------------------------------


def test_attachment_text_shows_content(tmp_path: Path) -> None:
    """Text attachment shows name and content in monospace."""
    att = Attachment(name="log.txt", mime_type="text/plain", text="hello world")
    step = _make_step(attachments=[att])
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "log.txt" in text
    assert "hello world" in text


def test_attachment_text_multiline(tmp_path: Path) -> None:
    """Multiline text attachment shows all lines."""
    att = Attachment(name="output.txt", mime_type="text/plain", text="line1\nline2\nline3")
    step = _make_step(attachments=[att])
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "line1" in text
    assert "line2" in text
    assert "line3" in text


def test_attachment_json_shows_content(tmp_path: Path) -> None:
    """JSON attachment shows name and content."""
    att = Attachment(
        name="data.json",
        mime_type="application/json",
        text='{"key": "value"}',
    )
    step = _make_step(attachments=[att])
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "data.json" in text
    assert '"key"' in text


def test_attachment_binary_shows_name_only(tmp_path: Path) -> None:
    """Binary attachment with no text shows name only."""
    att = Attachment(name="data.bin", mime_type="application/octet-stream", data_base64="AAAA")
    step = _make_step(attachments=[att])
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "data.bin" in text


def test_attachments_omitted_when_include_disabled(tmp_path: Path) -> None:
    """All attachments are omitted when include_attachments is False."""
    att = Attachment(
        name="screenshot.png",
        mime_type="image/png",
        data_base64="iVBORw0KGgo=",
    )
    step = _make_step(attachments=[att])
    run = _make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])])
    run_opts = ReportOptions(include_attachments=False)
    writer = DOCXWriter(run_opts)
    p = tmp_path / "report.docx"
    writer.write(run, str(p))
    text = _all_text(_read_docx(str(p)))
    assert "screenshot.png" not in text


def test_attachment_multiple_all_shown(tmp_path: Path) -> None:
    """Multiple attachments are all shown."""
    att1 = Attachment(name="a.txt", mime_type="text/plain", text="aaa")
    att2 = Attachment(name="b.json", mime_type="application/json", text='{"x":1}')
    att3 = Attachment(name="c.bin", mime_type="application/octet-stream")
    step = _make_step(attachments=[att1, att2, att3])
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "a.txt" in text
    assert "b.json" in text
    assert "c.bin" in text


def test_attachment_heading_present(tmp_path: Path) -> None:
    """Attachments section heading is present."""
    att = Attachment(name="log.txt", mime_type="text/plain", text="data")
    step = _make_step(attachments=[att])
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "Attachments" in text


def test_attachment_image_inserted_inline(tmp_path: Path) -> None:
    """Image attachment with valid PNG data is inserted inline."""

    png_b64 = (
        "iVBORw0KGgoAAAANSUhEUgAAAAoAAAAKCAIAAAACUFjqAAAAE0lEQVR4nGP8z4AP"
        "MOGVZRip0gBBLAETee26JgAAAABJRU5ErkJggg=="
    )
    att = Attachment(
        name="screenshot.png",
        mime_type="image/png",
        data_base64=png_b64,
    )
    step = _make_step(attachments=[att])
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    doc = _read_docx(path)
    assert len(doc.inline_shapes) >= 1


def test_attachment_image_invalid_data_falls_back(tmp_path: Path) -> None:
    """Image attachment with invalid data falls back to name listing."""
    att = Attachment(
        name="broken.png",
        mime_type="image/png",
        data_base64="!!!invalid!!!",
    )
    step = _make_step(attachments=[att])
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "broken.png" in text
    assert "could not display" in text


def test_attachment_text_empty_shows_name_only(tmp_path: Path) -> None:
    """Text attachment with empty text shows name only."""
    att = Attachment(name="empty.txt", mime_type="text/plain", text="")
    step = _make_step(attachments=[att])
    path = _write(_make_run([_make_feature(scenarios=[_make_scenario(steps=[step])])]), tmp_path)
    text = _all_text(_read_docx(path))
    assert "empty.txt" in text


# ---------------------------------------------------------------------------
# Branding: logo, primary_color
# ---------------------------------------------------------------------------


def test_docx_logo_on_cover(tmp_path: Path) -> None:
    """Logo is inserted on the cover page when logo_b64 is set."""

    png_b64 = (
        "iVBORw0KGgoAAAANSUhEUgAAAAoAAAAKCAIAAAACUFjqAAAAE0lEQVR4nGP8z4AP"
        "MOGVZRip0gBBLAETee26JgAAAABJRU5ErkJggg=="
    )
    opts = ReportOptions(logo_b64=f"data:image/png;base64,{png_b64}")
    run = _make_run([_make_feature()])
    writer = DOCXWriter(opts)
    p = tmp_path / "report.docx"
    writer.write(run, str(p))
    doc = _read_docx(str(p))
    # Logo image should be in the document
    assert len(doc.inline_shapes) >= 1


def test_docx_logo_invalid_data_skipped(tmp_path: Path) -> None:
    """Invalid logo data is silently skipped."""
    opts = ReportOptions(logo_b64="data:image/png;base64,!!!invalid!!!")
    run = _make_run([_make_feature()])
    writer = DOCXWriter(opts)
    p = tmp_path / "report.docx"
    writer.write(run, str(p))
    # Should still produce a valid docx with title
    text = _all_text(_read_docx(str(p)))
    assert "Test" in text


def test_docx_custom_primary_color(tmp_path: Path) -> None:
    """Custom primary color is applied to title."""
    opts = ReportOptions(primary_color="#1E90FF")
    run = _make_run([_make_feature()])
    writer = DOCXWriter(opts)
    p = tmp_path / "report.docx"
    writer.write(run, str(p))
    doc = _read_docx(str(p))
    # Title paragraph should have the custom color
    title_para = doc.paragraphs[0]
    title_run = title_para.runs[0]
    assert title_run.font.color.rgb == RGBColor(0x1E, 0x90, 0xFF)


def test_docx_default_primary_color(tmp_path: Path) -> None:
    """Default primary color is used when not specified."""
    opts = ReportOptions()
    run = _make_run([_make_feature()])
    writer = DOCXWriter(opts)
    p = tmp_path / "report.docx"
    writer.write(run, str(p))
    doc = _read_docx(str(p))
    title_para = doc.paragraphs[0]
    title_run = title_para.runs[0]
    assert title_run.font.color.rgb == RGBColor(0x25, 0x63, 0xEB)
