"""Tests for behave_modern_file_reports.docx_formatter."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest

from behave_modern_file_reports.docx_formatter import DOCXFormatter
from behave_modern_file_reports.models import (
    FeatureSummary,
    RunSummary,
    ScenarioResult,
    Step,
)

# ---------------------------------------------------------------------------
# Mock helpers
# ---------------------------------------------------------------------------


class MockStreamOpener:
    """Mock stream opener with a name attribute for path resolution."""

    def __init__(self, path: str) -> None:
        self.name = path
        self.opened = False

    def open(self) -> Any:
        self.opened = True
        return self


def _mock_config(userdata: dict[str, str] | None = None) -> SimpleNamespace:
    return SimpleNamespace(userdata=userdata or {})


def _make_simple_run() -> RunSummary:
    """Return a minimal run with one feature and one passed scenario."""
    step = Step(keyword="Given ", name="a step", status="passed", duration=0.01)
    scn = ScenarioResult(
        name="Test scenario",
        status="passed",
        duration=0.05,
        steps=[step],
    )
    feat = FeatureSummary(
        name="Test feature",
        duration=0.1,
        location="features/test.feature:1",
        scenarios=[scn],
    )
    return RunSummary(
        run_id="run_test",
        title="Test Report",
        features=[feat],
    )


def _read_docx(path: str) -> Any:
    from docx import Document
    return Document(path)


def _all_text(doc: Any) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    parts.append(para.text)
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Class attributes
# ---------------------------------------------------------------------------


def test_name_is_behave_modern_docx() -> None:
    assert DOCXFormatter.name == "behave-modern-docx"


def test_format_key_is_docx() -> None:
    assert DOCXFormatter._format_key == "docx"


def test_default_filename_is_report_docx() -> None:
    assert DOCXFormatter._default_filename == "report.docx"


def test_description_set() -> None:
    assert "docx" in DOCXFormatter.description.lower()


# ---------------------------------------------------------------------------
# Option resolution
# ---------------------------------------------------------------------------


def test_resolves_docx_specific_options() -> None:
    config = _mock_config({"bmfr.docx.title": "DOCX Title"})
    fmt = DOCXFormatter(config=config)
    assert fmt._options.title == "DOCX Title"


def test_docx_specific_overrides_global() -> None:
    config = _mock_config({
        "bmfr.title": "Global Title",
        "bmfr.docx.title": "DOCX Title",
    })
    fmt = DOCXFormatter(config=config)
    assert fmt._options.title == "DOCX Title"


def test_global_falls_back_for_docx() -> None:
    config = _mock_config({"bmfr.title": "Global Title"})
    fmt = DOCXFormatter(config=config)
    assert fmt._options.title == "Global Title"


# ---------------------------------------------------------------------------
# _write_report
# ---------------------------------------------------------------------------


def test_write_report_with_stream_opener(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    opener = MockStreamOpener(str(path))
    fmt = DOCXFormatter(stream_opener=opener)
    run = _make_simple_run()
    fmt._write_report(run, fmt._options)
    doc = _read_docx(str(path))
    text = _all_text(doc)
    assert "Test Report" in text
    assert "Test feature" in text
    assert "Test scenario" in text


def test_write_report_without_stream_opener(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    default_path = tmp_path / "report.docx"
    fmt = DOCXFormatter()
    monkeypatch.setattr(fmt, "_default_filename", str(default_path))
    run = _make_simple_run()
    fmt._write_report(run, fmt._options)
    doc = _read_docx(str(default_path))
    text = _all_text(doc)
    assert "Test Report" in text


def test_write_report_without_stream_opener_no_name(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    default_path = tmp_path / "report.docx"
    opener = SimpleNamespace()
    fmt = DOCXFormatter(stream_opener=opener)
    monkeypatch.setattr(fmt, "_default_filename", str(default_path))
    run = _make_simple_run()
    fmt._write_report(run, fmt._options)
    doc = _read_docx(str(default_path))
    text = _all_text(doc)
    assert "Test Report" in text


# ---------------------------------------------------------------------------
# close() lifecycle
# ---------------------------------------------------------------------------


def test_close_writes_report(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    opener = MockStreamOpener(str(path))
    fmt = DOCXFormatter(stream_opener=opener)
    fmt.feature(SimpleNamespace(
        name="F1", tags=[], location="f:1", description=None,
    ))
    fmt.scenario(SimpleNamespace(
        name="S1", tags=[], location="f:5",
        feature=SimpleNamespace(name="F1", tags=[], location=""),
        is_outline=False, rule=None, description=None,
    ))
    fmt.step(SimpleNamespace(
        keyword="Given ", name="step", status="passed",
        location="f:10", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.result(SimpleNamespace(
        keyword="Given ", name="step", status="passed",
        location="f:10", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.eof()
    fmt.close()
    doc = _read_docx(str(path))
    text = _all_text(doc)
    assert "F1" in text
    assert "S1" in text
    assert "step" in text


def test_close_is_idempotent(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    opener = MockStreamOpener(str(path))
    fmt = DOCXFormatter(stream_opener=opener)
    fmt.close()
    fmt.close()
    doc = _read_docx(str(path))
    text = _all_text(doc)
    assert text.count("Executive Summary") == 1


def test_close_without_stream_opener_writes_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    default_path = tmp_path / "report.docx"
    fmt = DOCXFormatter()
    monkeypatch.setattr(fmt, "_default_filename", str(default_path))
    fmt.close()
    doc = _read_docx(str(default_path))
    text = _all_text(doc)
    assert "Behave Modern Report" in text


# ---------------------------------------------------------------------------
# Visual design compliance
# ---------------------------------------------------------------------------


def test_heading1_for_feature(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    opener = MockStreamOpener(str(path))
    fmt = DOCXFormatter(stream_opener=opener)
    fmt.feature(SimpleNamespace(
        name="Login", tags=[], location="f:1", description=None,
    ))
    fmt.scenario(SimpleNamespace(
        name="Login ok", tags=[], location="f:5",
        feature=SimpleNamespace(name="Login", tags=[], location=""),
        is_outline=False, rule=None, description=None,
    ))
    fmt.step(SimpleNamespace(
        keyword="Given ", name="page", status="passed",
        location="f:10", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.result(SimpleNamespace(
        keyword="Given ", name="page", status="passed",
        location="f:10", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.eof()
    fmt.close()
    doc = _read_docx(str(path))
    headings = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert any("Login" in h.text for h in headings)


def test_heading2_for_scenario(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    opener = MockStreamOpener(str(path))
    fmt = DOCXFormatter(stream_opener=opener)
    fmt.feature(SimpleNamespace(
        name="F1", tags=[], location="f:1", description=None,
    ))
    fmt.scenario(SimpleNamespace(
        name="My Scenario", tags=[], location="f:5",
        feature=SimpleNamespace(name="F1", tags=[], location=""),
        is_outline=False, rule=None, description=None,
    ))
    fmt.step(SimpleNamespace(
        keyword="Given ", name="x", status="passed",
        location="f:10", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.result(SimpleNamespace(
        keyword="Given ", name="x", status="passed",
        location="f:10", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.eof()
    fmt.close()
    doc = _read_docx(str(path))
    headings = [p for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert any("My Scenario" in h.text for h in headings)


def test_status_badges_present(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    opener = MockStreamOpener(str(path))
    fmt = DOCXFormatter(stream_opener=opener)
    fmt.feature(SimpleNamespace(
        name="F1", tags=[], location="f:1", description=None,
    ))
    fmt.scenario(SimpleNamespace(
        name="S1", tags=[], location="f:5",
        feature=SimpleNamespace(name="F1", tags=[], location=""),
        is_outline=False, rule=None, description=None,
    ))
    fmt.step(SimpleNamespace(
        keyword="Given ", name="x", status="passed",
        location="f:10", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.result(SimpleNamespace(
        keyword="Given ", name="x", status="passed",
        location="f:10", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.eof()
    fmt.close()
    doc = _read_docx(str(path))
    text = _all_text(doc)
    assert "PASSED" in text


def test_step_tables_present(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    opener = MockStreamOpener(str(path))
    fmt = DOCXFormatter(stream_opener=opener)
    fmt.feature(SimpleNamespace(
        name="F1", tags=[], location="f:1", description=None,
    ))
    fmt.scenario(SimpleNamespace(
        name="S1", tags=[], location="f:5",
        feature=SimpleNamespace(name="F1", tags=[], location=""),
        is_outline=False, rule=None, description=None,
    ))
    fmt.step(SimpleNamespace(
        keyword="Given ", name="step1", status="passed",
        location="f:10", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.result(SimpleNamespace(
        keyword="Given ", name="step1", status="passed",
        location="f:10", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.eof()
    fmt.close()
    doc = _read_docx(str(path))
    text = _all_text(doc)
    assert "step1" in text
    assert "Status" in text


def test_error_block_present(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    opener = MockStreamOpener(str(path))
    fmt = DOCXFormatter(stream_opener=opener)
    fmt.feature(SimpleNamespace(
        name="F1", tags=[], location="f:1", description=None,
    ))
    fmt.scenario(SimpleNamespace(
        name="S1", tags=[], location="f:5",
        feature=SimpleNamespace(name="F1", tags=[], location=""),
        is_outline=False, rule=None, description=None,
    ))
    try:
        raise AssertionError("fail!") from None
    except AssertionError as exc:
        failed_step = SimpleNamespace(
            keyword="Then ", name="bad", status="failed",
            location="f:15", duration=0.01, text=None,
            error=exc, exception=None, error_message=None,
        )
    fmt.step(SimpleNamespace(
        keyword="Then ", name="bad", status="failed",
        location="f:15", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.result(failed_step)
    fmt.eof()
    fmt.close()
    doc = _read_docx(str(path))
    text = _all_text(doc)
    assert "Failure" in text
    assert "fail!" in text


# ---------------------------------------------------------------------------
# Full lifecycle with options
# ---------------------------------------------------------------------------


def test_full_lifecycle_with_options(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    opener = MockStreamOpener(str(path))
    config = _mock_config({"bmfr.title": "Custom DOCX Report"})
    fmt = DOCXFormatter(stream_opener=opener, config=config)
    assert fmt._options.title == "Custom DOCX Report"

    fmt.feature(SimpleNamespace(
        name="Auth", tags=["security"], location="f:1",
        description="Auth flows.",
    ))
    fmt.scenario(SimpleNamespace(
        name="Login ok", tags=["smoke"], location="f:5",
        feature=SimpleNamespace(name="Auth", tags=[], location=""),
        is_outline=False, rule=None, description=None,
    ))
    fmt.step(SimpleNamespace(
        keyword="Given ", name="user exists", status="passed",
        location="f:10", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.result(SimpleNamespace(
        keyword="Given ", name="user exists", status="passed",
        location="f:10", duration=0.01, text=None,
        error=None, exception=None, error_message=None,
    ))
    fmt.eof()
    fmt.close()
    doc = _read_docx(str(path))
    text = _all_text(doc)
    assert "Custom DOCX Report" in text
    assert "Auth" in text
    assert "Login ok" in text
    assert "user exists" in text
    assert "security" in text
    assert "smoke" in text


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def test_entry_point_resolvable() -> None:
    try:
        from importlib.metadata import entry_points
    except ImportError:
        from importlib_metadata import entry_points  # type: ignore

    eps = entry_points()
    if hasattr(eps, "select"):
        formatter_eps = eps.select(group="behave.formatters")
    else:
        formatter_eps = eps.get("behave.formatters", [])  # type: ignore
    names = [ep.name for ep in formatter_eps]
    assert "behave-modern-docx" in names
